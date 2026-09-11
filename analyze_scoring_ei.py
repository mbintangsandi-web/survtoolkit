"""
analyze_scoring_ei.py
=====================
Deteksi Error Input (EI) untuk QA Scorecard NOC Surveillance
Kategori: AM/PM Error + Submit Before Event
Penalti  : -1 per tiket (satu penalti meskipun ada multi-error)
Beban    : Inputer yang bertugas di shift saat kejadian (dari roster Excel)

CARA PAKAI:
    python analyze_scoring_ei.py --csv "Form Surveillance NOC.csv" --xlsx "JADWAL_PIKET_NOC_SURV.xlsx" --out ei_output.json

Kemudian copy isi ei_output.json ke scoring.html:
  - Tambahkan field "EI" ke tiap entri monthlyAgg
  - Tambahkan key "ei" ke detailLists
"""

import argparse, json, re
from datetime import datetime, timedelta
from collections import defaultdict

import pandas as pd
from openpyxl import load_workbook

try:
    from rapidfuzz import fuzz
    HAS_FUZZY = True
except ImportError:
    HAS_FUZZY = False
    print("[WARN] rapidfuzz tidak terinstall. Gunakan: pip install rapidfuzz")

# -----------------------------------------------------------------------
# KONFIGURASI
# -----------------------------------------------------------------------
NAMA_RESMI = [
    "IHSANUL MAULANA", "MUHAMMAD FARID", "MUHAMMAD HELMI",
    "ALFITO SRI PANGESTU", "EDI GUNTORO", "BAYU RAIHAN SANJAYA",
    "ADE PRASETYO", "MUHAMMAD BINTANG SANDI JANNATA",
]
BULAN_FILTER  = {5: "Mei", 6: "Juni", 7: "Juli", 8: "Agustus"}
HANDOVER_MIN  = 30       # menit toleransi handover
AMPM_MIN_H    = 11.0     # jam minimum selisih agar dianggap AM/PM error
AMPM_MAX_H    = 13.0     # jam maksimum
FUZZY_THRESH  = 75       # threshold kecocokan nama (0-100)

DT_FORMATS = [
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
    "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M",
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
]

# -----------------------------------------------------------------------
# UTILITAS
# -----------------------------------------------------------------------
def parse_dt(s):
    if not s or (isinstance(s, float) and str(s) == "nan"):
        return None
    s = str(s).strip()
    for fmt in DT_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None

def norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip().upper())

def shift_of(dt):
    h = dt.hour
    if 7 <= h < 15: return "P"
    if 15 <= h < 23: return "S"
    return "M"

def fuzzy_best(raw, candidates=NAMA_RESMI):
    if not HAS_FUZZY:
        raw_u = norm(raw)
        for c in candidates:
            if raw_u == c: return c, 100
        return None, 0
    best, score = None, 0
    for c in candidates:
        s = fuzz.token_set_ratio(raw, c)
        if s > score:
            score, best = s, c
    return (best, score) if score >= FUZZY_THRESH else (None, score)

def detect_ei(ts_dt, down_dt):
    """Return 'ampm' | 'submit_before' | None. Prioritas: ampm > submit_before."""
    if ts_dt is None or down_dt is None:
        return None
    delta_h   = abs((ts_dt - down_dt).total_seconds()) / 3600
    delta_sec = (ts_dt - down_dt).total_seconds()
    if AMPM_MIN_H <= delta_h <= AMPM_MAX_H:
        return "ampm"
    if delta_sec < -(HANDOVER_MIN * 60):
        return "submit_before"
    return None

# -----------------------------------------------------------------------
# LOAD ROSTER
# -----------------------------------------------------------------------
def load_roster(xlsx_path):
    """Return { (date, shift_code) -> [nama_resmi, ...] }"""
    roster = {}
    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        print(f"[WARN] Gagal buka Excel: {e}")
        return roster

    for sheet_name in wb.sheetnames:
        if "SEPTEMBER" in sheet_name.upper():
            continue
        if not any(k in sheet_name.upper() for k in ["LINE", "SUMMARY", "LINEUP"]):
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # Cari baris header
        header, hidx = None, 0
        for i, row in enumerate(rows):
            ru = [str(c).upper().strip() if c else "" for c in row]
            if any("TANGGAL" in c or "DATE" in c for c in ru):
                header, hidx = ru, i
                break
        if header is None:
            continue

        cm = {}
        for j, h in enumerate(header):
            if "TANGGAL" in h or "DATE" in h: cm["date"] = j
            elif h in ("P", "PAGI"):           cm["P"] = j
            elif h in ("S", "SIANG"):          cm["S"] = j
            elif h in ("M", "MALAM"):          cm["M"] = j

        if "date" not in cm:
            continue

        for row in rows[hidx + 1:]:
            dv = row[cm["date"]]
            if dv is None:
                continue
            try:
                d = dv.date() if isinstance(dv, datetime) else pd.to_datetime(str(dv), dayfirst=True).date()
            except Exception:
                continue
            for sc in ("P", "S", "M"):
                if sc not in cm or not row[cm[sc]]:
                    continue
                parts = re.split(r"[,\n/]", str(row[cm[sc]]))
                resolved = []
                for raw_name in parts:
                    n = norm(raw_name)
                    if not n:
                        continue
                    matched, _ = fuzzy_best(n)
                    if matched:
                        resolved.append(matched)
                if resolved:
                    roster[(d, sc)] = resolved
    return roster

def duty_person(down_dt, roster):
    sc = shift_of(down_dt)
    d  = down_dt.date()
    for candidate_date in [d, d - timedelta(days=1)]:
        names = roster.get((candidate_date, sc), [])
        if names:
            return names[0]
    return None

# -----------------------------------------------------------------------
# ANALISIS
# -----------------------------------------------------------------------
def analyze(csv_path, xlsx_path):
    print(f"[INFO] Baca CSV: {csv_path}")
    df = pd.read_csv(csv_path, dtype=str)
    df.columns = [c.strip() for c in df.columns]
    col_upper  = {c.upper(): c for c in df.columns}

    def gc(*cands):
        for c in cands:
            if c.upper() in col_upper: return col_upper[c.upper()]
        return None

    COL_TS   = gc("Timestamp", "TIMESTAMP")
    COL_DOWN = gc("DEVICE DOWN TIME", "DOWN TIME", "DOWNTIME")
    COL_PIC  = gc("PIC BLAST DOWN", "PIC DOWN")

    if not all([COL_TS, COL_DOWN, COL_PIC]):
        print(f"[ERROR] Kolom wajib tidak ditemukan. Kolom tersedia: {list(df.columns)}")
        return {}, {}

    roster = load_roster(xlsx_path)
    print(f"[INFO] Entri roster: {len(roster)}")

    ei_counts = defaultdict(lambda: defaultdict(int))
    ei_detail = defaultdict(list)
    detected = 0

    for _, row in df.iterrows():
        ts_dt   = parse_dt(row.get(COL_TS))
        down_dt = parse_dt(row.get(COL_DOWN))
        if ts_dt is None or down_dt is None:
            continue
        if down_dt.month not in BULAN_FILTER:
            continue

        ei_type = detect_ei(ts_dt, down_dt)
        if ei_type is None:
            continue

        detected += 1
        bulan = BULAN_FILTER[down_dt.month]

        person = duty_person(down_dt, roster)
        if person is None:
            person, _ = fuzzy_best(norm(row.get(COL_PIC, "")))
        if person is None:
            continue

        ei_counts[person][bulan] += 1
        ticket_id = str(row.get("ID", row.get("TICKET ID", row.get("NO TIKET", "")))).strip()
        ei_detail[person].append({
            "id"       : ticket_id or "-",
            "down"     : str(row.get(COL_DOWN, "-")).strip(),
            "submit"   : str(row.get(COL_TS, "-")).strip(),
            "delta_jam": round(abs((ts_dt - down_dt).total_seconds()) / 3600, 1),
            "jenis"    : "AM/PM Error" if ei_type == "ampm" else "Submit Sebelum Event",
            "regional" : str(row.get(gc("REGIONAL", "REGION") or "", "-")).strip(),
            "hostname" : str(row.get(gc("HOSTNAME", "DEVICE NAME") or "", "-")).strip(),
            "bulan"    : bulan,
        })

    print(f"[INFO] Total EI terdeteksi: {detected}")
    return ei_counts, ei_detail

# -----------------------------------------------------------------------
# OUTPUT
# -----------------------------------------------------------------------
def generate_output(ei_counts, ei_detail):
    bulan_list = list(BULAN_FILTER.values())
    by_month   = {b: {} for b in bulan_list + ["Semua (Mei-Agustus)"]}

    for nama in NAMA_RESMI:
        total = 0
        for b in bulan_list:
            v = ei_counts.get(nama, {}).get(b, 0)
            by_month[b][nama] = v
            total += v
        by_month["Semua (Mei-Agustus)"][nama] = total

    return {"ei_by_month": by_month, "ei_detail": dict(ei_detail)}

# -----------------------------------------------------------------------
# MAIN
# -----------------------------------------------------------------------
def main():
    p = argparse.ArgumentParser()
    p.add_argument("--csv",  required=True)
    p.add_argument("--xlsx", required=True)
    p.add_argument("--out",  default="ei_output.json")
    args = p.parse_args()

    counts, detail = analyze(args.csv, args.xlsx)
    out = generate_output(counts, detail)

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    print(f"\n[DONE] Output: {args.out}")
    print("\n=== Ringkasan EI (Semua Bulan) ===")
    for nama in NAMA_RESMI:
        n = out["ei_by_month"]["Semua (Mei-Agustus)"].get(nama, 0)
        print(f"  {nama:<40}: {n} tiket EI")

if __name__ == "__main__":
    main()
